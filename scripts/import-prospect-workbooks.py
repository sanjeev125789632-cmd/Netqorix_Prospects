"""Regenerate the five added lead batches from the supplied XLSX workbooks.

Usage: python scripts/import-prospect-workbooks.py /path/to/workbooks
Requires openpyxl. The four original batches retain their existing generated files
and IDs so browser-local tracking survives this import.
"""

import json
import sys
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parent.parent
SOURCES = [
    ("Netqorix_Bankura_Durgapur_Prospects.xlsx", "Leads", "Bankura & Durgapur", "bankura", "2026-09-23"),
    ("Netqorix_JammuKashmir_Prospects.xlsx", "Leads", "Jammu & Kashmir", "jammu-kashmir", "2026-09-23"),
    ("Netqorix_NorthEast_Prospects_1.xlsx", "Leads", "North East", "north-east", "2026-09-21"),
    ("Netqorix_International_Round5.xlsx", "Overseas Leads", "International Round 5", "international-5", "2026-09-13"),
    ("Netqorix_Round6_National_and_International.xlsx", "All Leads", None, "round-6", "2026-09-18"),
]


def text(value):
    return str(value).strip() if value is not None else ""


def number(value):
    try:
        return float(value or 0)
    except (TypeError, ValueError):
        return 0


def clean_email(value):
    value = text(value)
    return value if "@" in value and not value.startswith("—") else ""


def import_source(directory, filename, sheet, region, slug, date):
    # Normal mode retains hyperlink targets. In read-only mode these cells only
    # expose the display text "Open listing" and silently lose their URLs.
    workbook = load_workbook(directory / filename, read_only=False, data_only=True)
    worksheet = workbook[sheet]
    rows = worksheet.values
    headers = [text(cell) for cell in next(rows)]
    result = []
    maps_col = headers.index("Maps listing") + 1
    for row_num, cells in enumerate(rows, start=2):
        row = dict(zip(headers, cells))
        if not isinstance(row.get("Rank"), int) or not row.get("Business"):
            continue
        scope = text(row.get("Scope"))
        area = region or ("Round 6 International" if scope == "International" else "Round 6 National")
        city = text(row.get("Town") or row.get("City"))
        locality = text(row.get("Locality") or row.get("Address")) or city
        rank = row["Rank"]
        lead = {
            "id": f"lead-{slug}-{scope.lower()}-{rank}" if scope else f"lead-{slug}-{rank}",
            "region": area,
            "rank": rank,
            "tier": text(row.get("Tier")),
            "fit": int(number(row.get("Fit /100"))),
            "business": text(row.get("Business")),
            "category": text(row.get("Category")),
            "segment": text(row.get("Segment")),
            "locality": locality,
            "city": city,
            "phone": text(row.get("Phone")) or "— none listed",
            "rating": number(row.get("Rating")),
            "reviews": int(number(row.get("Reviews"))),
            "package": text(row.get("Package")),
            "listPrice": text(row.get("List price")),
            "dealValue": number(row.get("Deal value (₹)")),
            "bestCallWindow": text(row.get("Best call window (IST)")),
            "pitchAngle": text(row.get("Pitch angle")),
            "intent": int(number(row.get("Intent"))),
            "volume": int(number(row.get("Volume"))),
            "ratingPts": int(number(row.get("Rating pts"))),
            "reach": int(number(row.get("Reach"))),
            "mapsUrl": text(worksheet.cell(row_num, maps_col).hyperlink.target if worksheet.cell(row_num, maps_col).hyperlink else row.get("Maps listing")),
            "sourceFile": filename,
            "sourceDate": date,
        }
        if row.get("Address") and text(row["Address"]) not in ("— not published", "—"):
            lead["address"] = text(row["Address"])
        if row.get("State"):
            lead["market"] = text(row["State"])
        if row.get("Market"):
            lead["market"] = text(row["Market"])
        if row.get("Country"):
            lead["country"] = text(row["Country"])
        elif area in ("International Round 5", "Round 6 International"):
            lead["country"] = text(row.get("Market"))
        else:
            lead["country"] = "IN"
        email = clean_email(row.get("Email"))
        if email:
            lead["sourceEmail"] = email
        result.append(lead)
    workbook.close()
    return result


def main(directory):
    all_leads = []
    for filename, sheet, region, slug, date in SOURCES:
        leads = import_source(directory, filename, sheet, region, slug, date)
        print(f"{filename}: {len(leads)}")
        all_leads.extend(leads)
    ids = [lead["id"] for lead in all_leads]
    if len(all_leads) != 984 or len(set(ids)) != len(ids):
        raise ValueError("Expected 984 new leads with unique IDs")
    output = ROOT / "src/data/prospectsAdditional.ts"
    output.write_text(
        "import type { Prospect } from '../types/prospect';\n\n"
        "// Generated from the five September 2026 workbooks by scripts/import-prospect-workbooks.py.\n"
        "export const prospectsAdditional: Prospect[] = "
        + json.dumps(all_leads, ensure_ascii=False, separators=(",", ":"))
        + ";\n",
        encoding="utf-8",
    )
    print(f"Wrote {output}: {len(all_leads)} new leads")


if __name__ == "__main__":
    main(Path(sys.argv[1] if len(sys.argv) > 1 else "."))
